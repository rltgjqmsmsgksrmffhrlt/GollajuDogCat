"""리뷰 설문 답변과 본문을 실제 입력 폼(기호성만 필수, 나머지 선택) 기준으로 현실화한다.

- PALATABILITY: 필수 -> 항상 응답.  DIGESTION/SKIN_COAT/WEIGHT_VITALITY/ALLERGY/FEEDING_CONVENIENCE: 선택
  -> 확률적으로 미응답(review_answer = NULL, 행은 유지). 건강 고민 태그와 관련된 문항·사용기간이 긴 리뷰는 응답률 상향.
- 4~5점 리뷰(235건): 응답한 문항의 답변을 현실적인 비율로 다시 뽑고, 본문을 응답 내용에 맞춰 재작성
  (본문에 언급되는 항목은 반드시 응답이 있는 항목).
- 1~3점 리뷰(52건): 본문과 주 불만 항목 답변은 유지, 그 외 선택 문항만 미응답 처리 가능.
- 답변 매핑: NEGATIVE=나빠졌어요/안 먹어요, NEUTRAL=그대로예요/보통이에요, POSITIVE=좋아졌어요/잘 먹어요
시드 고정(2028). 이미 적용돼 있으면 종료.
"""
import os
import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
PATH = os.path.join(HERE, "..", "data", "mock", "상품 목데이터",
                    "golajulgaenyang_full_schema_mockdata_product_thumbnails_linked.xlsx")
RNG = np.random.default_rng(2028)
MARK = "리뷰 설문 응답"

# 1~3점 본문에서 주 불만 항목을 찾기 위해 별점 조정 스크립트의 문구 풀을 재사용
_src = open(os.path.join(HERE, "mock_rebalance_review_ratings.py"), encoding="utf-8").read()
_ns = {"__file__": os.path.join(HERE, "mock_rebalance_review_ratings.py")}
exec(_src.split("wb = openpyxl.load_workbook(PATH)")[0], _ns)
LOW_POOLS = {"PALATABILITY": [p for v in list(_ns["PAL"].values()) + list(_ns["PAL_MILD"].values()) for p in v]}
for k in _ns["ISSUE_TXT"]:
    LOW_POOLS[k] = _ns["ISSUE_TXT"][k] + _ns["ISSUE_MILD"][k]


def main_issue(text):
    for k, pool in LOW_POOLS.items():
        if any(p in text for p in pool):
            return k
    return None


OPTIONAL = ["DIGESTION", "SKIN_COAT", "WEIGHT_VITALITY", "ALLERGY", "FEEDING_CONVENIENCE"]
FILL = {"DIGESTION": .65, "SKIN_COAT": .50, "WEIGHT_VITALITY": .45, "ALLERGY": .35, "FEEDING_CONVENIENCE": .40}
CONCERN = {"DIGESTION": {"DIGESTION"}, "SKIN_COAT": {"SKIN_COAT", "HAIRBALL", "EYE"},
           "WEIGHT_VITALITY": {"WEIGHT", "JOINT"}, "ALLERGY": {"SKIN_COAT"}, "FEEDING_CONVENIENCE": set()}
SLOW = {"SKIN_COAT", "WEIGHT_VITALITY"}  # 효과 확인에 시간이 걸리는 문항
# 4~5점 리뷰의 응답 분포 (POSITIVE, NEUTRAL)
P_POS = {"DIGESTION": .55, "SKIN_COAT": .45, "WEIGHT_VITALITY": .40, "ALLERGY": .08, "FEEDING_CONVENIENCE": .60}

MENTION = {
    ("DIGESTION", "POSITIVE"): ["급여 후 변 상태가 좋아졌어요", "배변이 한결 좋아졌어요", "변이 단단해지고 냄새도 덜해요"],
    ("DIGESTION", "NEUTRAL"): ["배변 상태는 그대로예요", "배변은 크게 달라지지 않았어요", "변 상태는 평소와 비슷해요"],
    ("SKIN_COAT", "POSITIVE"): ["털에 윤기가 생겼어요", "긁는 횟수가 줄고 모질이 좋아졌어요", "피부 상태가 한결 나아졌어요"],
    ("SKIN_COAT", "NEUTRAL"): ["피부와 털 상태는 그대로예요", "모질은 아직 큰 변화는 없어요", "털 상태는 비슷해요"],
    ("WEIGHT_VITALITY", "POSITIVE"): ["활력이 좋아졌어요", "산책할 때 기운이 더 넘쳐요", "체중이 안정되고 활동량이 늘었어요"],
    ("WEIGHT_VITALITY", "NEUTRAL"): ["체중과 활력은 그대로예요", "체중 변화는 아직 없어요", "활동량은 평소와 비슷해요"],
    ("ALLERGY", "POSITIVE"): ["가려움이 줄어든 느낌이에요", "알러지 증상이 한결 완화됐어요"],
    ("ALLERGY", "NEUTRAL"): ["알러지 반응은 없었어요", "먹고 나서 이상 반응은 없었어요", "알러지 증상 없이 잘 먹고 있어요"],
    ("FEEDING_CONVENIENCE", "POSITIVE"): ["급여하기 편해요", "소분이 쉬워서 관리하기 좋아요", "한 번에 먹이기 편한 형태예요"],
    ("FEEDING_CONVENIENCE", "NEUTRAL"): ["급여 방법은 무난해요", "급여하기에 보통 정도예요"],
}
OPEN_POS = ["{s} 처음부터 잘 먹었어요.", "{s} 잘 먹어서 만족해요.", "{n} 기호성이 좋아요."]
OPEN_NEU = ["{n} 반응은 보통이에요.", "{s} 무난하게 먹어요.", "{t} 먹긴 하는데 엄청 좋아하진 않아요."]
CLOSE_POS = ["재구매하려고 합니다.", "다음에도 구매할 예정이에요.", "전체적으로 만족해서 재구매 의사 있어요."]
CLOSE_NEU = ["조금 더 급여해보고 재구매 결정할 것 같아요.", "좀 더 지켜볼게요."]


def josa(name, with_final, without_final):
    """이름 끝 글자 받침 유무에 따라 조사를 붙인다."""
    return name + (with_final if (ord(name[-1]) - 0xAC00) % 28 else without_final)


def pick(lst):
    return lst[RNG.integers(len(lst))]


wb = openpyxl.load_workbook(PATH)
rs = wb["README"]
if any(str(c.value).startswith(MARK) for c in rs["A"]):
    print("이미 적용돼 있습니다."); raise SystemExit

def tbl(ws):
    h = {c.value: c.column - 1 for c in ws[1]}
    return h, [r for r in ws.iter_rows(min_row=2) if r[0].value is not None]

rh, rrows = tbl(wb["review"])
qh, qrows = tbl(wb["review_question"])
ch, crows = tbl(wb["review_health_concern"])
ph, prows = tbl(wb["review_pet"])
pet = {r[ph["review_id"]].value: r[ph["name"]].value for r in prows}
concern = {}
for r in crows:
    concern.setdefault(r[ch["review_id"]].value, set()).add(r[ch["concern_code"]].value)
q = {(r[qh["review_id"]].value, r[qh["review_question_type"]].value): r[qh["review_answer"]] for r in qrows}

stat = {t: [0, 0] for t in OPTIONAL}
for r in rrows:
    rid, star = r[rh["id"]].value, r[rh["star_rate"]].value
    usage, tags = r[rh["usage_period"]].value, concern.get(rid, set())
    name = pet[rid]
    low_main = main_issue(r[rh["text"]].value) if star <= 3 else None
    if star <= 3 and r[rh["text"]].value.startswith(name + "이 "):  # "루이이" -> "루이가" 조사 보정
        r[rh["text"]].value = josa(name, "이", "가") + r[rh["text"]].value[len(name) + 1:]
    answered = {}
    for t in OPTIONAL:
        related = bool(tags & CONCERN[t])
        p = .9 if related else FILL[t] * (.5 if (t in SLOW and usage < 14) else 1)
        keep = RNG.random() < p or t == low_main
        cell = q[(rid, t)]
        if star >= 4:
            if keep:
                pp = P_POS[t] + (.25 if related and t != "ALLERGY" else 0)
                cell.value = "POSITIVE" if RNG.random() < pp else "NEUTRAL"
        if not keep:
            cell.value = None
        if cell.value is not None:
            answered[t] = cell.value
        stat[t][0] += 1; stat[t][1] += cell.value is not None
    if star >= 4:  # 본문 재작성
        pal = q[(rid, "PALATABILITY")].value
        pos = pal == "POSITIVE"
        parts = [pick(OPEN_POS if pos else OPEN_NEU).format(n=name, s=josa(name, "이", "가"), t=josa(name, "은", "는"))]
        if usage >= 30 and RNG.random() < .4:
            parts.append(f"{usage}일 정도 먹여봤어요.")
        order = list(RNG.permutation(OPTIONAL))
        for t in order:
            if t in answered and answered[t] in ("POSITIVE", "NEUTRAL"):
                related = bool(tags & CONCERN[t])
                if RNG.random() < (.9 if related else .55):
                    parts.append(pick(MENTION[(t, answered[t])]) + ".")
        parts.append(pick(CLOSE_POS if (pos and star == 5) else CLOSE_NEU if not pos else CLOSE_POS + CLOSE_NEU))
        text = " ".join(parts)
        assert len(text) <= 255
        r[rh["text"]].value = text

rs.cell(rs.max_row + 1, 1, MARK)
rs.cell(rs.max_row, 2, "기호성(PALATABILITY)만 필수 응답, 나머지 5개 문항은 선택이라 일부 리뷰는 review_answer NULL "
        "(리뷰당 6행은 유지). 본문에 언급된 항목은 모두 응답이 있는 항목")
wb.save(PATH)
for t, (n, a) in stat.items():
    print(f"{t:22s} 응답률 {a / n:.0%}")
