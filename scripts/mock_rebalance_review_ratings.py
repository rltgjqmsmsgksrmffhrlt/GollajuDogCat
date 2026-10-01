"""상품 목데이터 리뷰 별점을 현실적인 J자형 분포로 재조정한다.

대상: data/mock/상품 목데이터/golajulgaenyang_full_schema_mockdata_product_thumbnails_linked.xlsx
- 별점 목표(287건): 5점 180 / 4점 55 / 3점 18 / 2점 10 / 1점 24 (평균 약 4.24)
- 낮은 별점으로 바뀐 리뷰는 review.text, review_question.review_answer 도 별점에 맞게 재작성
- products.avg_rating, validation 시트, README 를 실제 집계에 맞게 갱신
시드 고정(2026) -> 재실행해도 동일 결과. 이미 조정된 파일에 다시 돌리면 변경 없음(별점 분포 확인 후 종료).
"""
import os
import numpy as np
import openpyxl

DIR = os.path.join(os.path.dirname(__file__), "..", "data", "mock", "상품 목데이터")
PATH = os.path.join(DIR, "golajulgaenyang_full_schema_mockdata_product_thumbnails_linked.xlsx")
RNG = np.random.default_rng(2026)

# (원래 별점 -> 새 별점, 건수)
PLAN = [(4, 3, 18), (4, 2, 1), (5, 2, 9), (5, 1, 24)]

ISSUES = ["PALATABILITY", "DIGESTION", "SKIN_COAT", "ALLERGY", "FEEDING_CONVENIENCE"]

PAL = {
    "FOOD": ["사료를 거의 먹지 않고 남겼어요", "냄새를 맡고 고개를 돌려요", "기존 사료보다 먹는 양이 확 줄었어요"],
    "TREAT": ["간식인데도 관심을 보이지 않아요", "한두 번 먹고 더는 안 먹으려 해요", "기대한 만큼 잘 먹지 않아요"],
    "SUPPLEMENT": ["영양제를 먹이려고 하면 피해요", "섞어줘도 알아채고 안 먹어요", "급여 자체가 힘들어요"],
}
ISSUE_TXT = {
    "DIGESTION": ["급여 후 묽은 변을 봤어요", "먹고 나서 배변 상태가 안 좋아졌어요", "구토 증상이 있어서 급여를 멈췄어요"],
    "SKIN_COAT": ["급여 후 긁는 횟수가 늘었어요", "피부가 붉어지고 털이 푸석해졌어요", "눈에 띄는 모질 개선이 없었어요"],
    "ALLERGY": ["알레르기 반응이 의심돼서 중단했어요", "먹고 나서 귀와 발을 계속 핥았어요", "기존에 없던 가려움이 생겼어요"],
    "FEEDING_CONVENIENCE": ["포장이 불편해서 급여하기 어려워요", "소분이 어렵고 흘리기 쉬워요", "알갱이 크기가 맞지 않아 먹기 불편해해요"],
}
# 3점용 순한 표현 (풀 크기는 위와 같게 유지해 난수 순서를 보존)
PAL_MILD = {
    "FOOD": ["처음보다 먹는 양이 조금 줄었어요", "잘 먹을 때도 있고 남길 때도 있어요", "기호성은 그저 그런 편이에요"],
    "TREAT": ["간식 반응이 그리 크지 않았어요", "좋아하는 날도 있고 시큰둥한 날도 있어요", "기대만큼 열광하지는 않아요"],
    "SUPPLEMENT": ["먹이는 데 약간 손이 가요", "간식에 섞으면 먹지만 그냥은 잘 안 먹어요", "급여 반응이 들쭉날쭉해요"],
}
ISSUE_MILD = {
    "DIGESTION": ["변 상태가 가끔 무르긴 했어요", "배변 변화는 뚜렷하지 않았어요", "처음 며칠은 배변이 살짝 달라졌어요"],
    "SKIN_COAT": ["피부나 모질 변화는 아직 잘 모르겠어요", "눈에 띄는 모질 개선은 못 느꼈어요", "긁는 횟수는 비슷했어요"],
    "ALLERGY": ["알레르기 반응은 없었지만 큰 변화도 없었어요", "특별한 이상은 없었는데 만족스럽진 않아요", "가려움은 비슷한 수준이에요"],
    "FEEDING_CONVENIENCE": ["포장이 조금 불편했어요", "소분이 약간 번거로워요", "알갱이 크기가 조금 아쉬워요"],
}
OPEN = {1: ["{n}에게는 맞지 않았어요.", "{n} 급여 후 문제가 있었어요.", "{n}에게 권하고 싶지 않아요."],
        2: ["{n}에게는 아쉬운 제품이었어요.", "{n} 반응이 기대보다 별로였어요."],
        3: ["{n} 반응은 보통이에요.", "{n}에게는 그냥 무난한 정도였어요."]}
CLOSE = {1: "재구매 의사 없습니다.", 2: "재구매는 고민 중이에요.", 3: "조금 더 먹여보고 결정할 것 같아요."}


def make_text(name, star, issue, cat):
    if star == 3:
        pool = PAL_MILD[cat] if issue == "PALATABILITY" else ISSUE_MILD[issue]
    else:
        pool = PAL[cat] if issue == "PALATABILITY" else ISSUE_TXT[issue]
    o = OPEN[star][RNG.integers(len(OPEN[star]))].format(n=name)
    b = pool[RNG.integers(len(pool))] + "."
    if star == 3:
        b = "다만 " + b
    return f"{o} {b} {CLOSE[star]}"


wb = openpyxl.load_workbook(PATH)


def table(ws, hr=1):
    hdr = {c.value: c.column for c in ws[hr]}
    first = min(hdr.values()) - 1
    return hdr, [r for r in ws.iter_rows(min_row=hr + 1) if r[first].value is not None]


rv_ws = wb["review"]; rh, rrows = table(rv_ws)
cur = {}
for r in rrows:
    cur[r[rh["star_rate"] - 1].value] = cur.get(r[rh["star_rate"] - 1].value, 0) + 1
if any(s < 4 for s in cur):
    print("이미 조정된 파일입니다:", dict(sorted(cur.items()))); raise SystemExit

prod_ws = wb["products"]; ph, prows = table(prod_ws)
cat = {r[ph["id"] - 1].value: r[ph["category_code"] - 1].value for r in prows}
pet_ws = wb["review_pet"]; peh, petrows = table(pet_ws)
pet = {r[peh["review_id"] - 1].value: r[peh["name"] - 1].value for r in petrows}
q_ws = wb["review_question"]; qh, qrows = table(q_ws)
qidx = {(r[qh["review_id"] - 1].value, r[qh["review_question_type"] - 1].value): r for r in qrows}

by_star = {s: [r for r in rrows if r[rh["star_rate"] - 1].value == s] for s in (4, 5)}
chosen = {4: iter(RNG.permutation(len(by_star[4]))), 5: iter(RNG.permutation(len(by_star[5])))}
for old, new, n in PLAN:
    for _ in range(n):
        r = by_star[old][next(chosen[old])]
        rid = r[rh["id"] - 1].value
        pc = cat[r[rh["product_id"] - 1].value]
        issues = list(RNG.permutation(ISSUES))
        # 기호성은 식품/간식/영양제 모두 가장 흔한 불만 -> 우선 배치
        main = "PALATABILITY" if RNG.random() < 0.5 else issues[0]
        r[rh["star_rate"] - 1].value = new
        r[rh["text"] - 1].value = make_text(pet[rid], new, main, pc)
        # 설문 답변: 3점=주 이슈 NEUTRAL, 2점=주 이슈 NEGATIVE, 1점=주 이슈+1개 NEGATIVE
        neg = [] if new == 3 else [main] + ([i for i in issues if i != main][:1] if new == 1 else [])
        for t in ISSUES:
            ans = "NEGATIVE" if t in neg else ("NEUTRAL" if (t == main or new < 3) else qidx[(rid, t)][qh["review_answer"] - 1].value)
            qidx[(rid, t)][qh["review_answer"] - 1].value = ans

# 상품별 집계 재계산
agg = {}
for r in rrows:
    if r[rh["deleted_at"] - 1].value is None:
        agg.setdefault(r[rh["product_id"] - 1].value, []).append(r[rh["star_rate"] - 1].value)
avg = {k: round(sum(v) / len(v), 1) for k, v in agg.items()}
for r in prows:
    r[ph["avg_rating"] - 1].value = avg[r[ph["id"] - 1].value]
    assert r[ph["review_count"] - 1].value == len(agg[r[ph["id"] - 1].value])
vh, vrows = table(wb["validation"])
for r in vrows:
    a = avg[r[vh["product_id"] - 1].value]
    r[vh["product.avg_rating"] - 1].value = a
    r[vh["actual_avg_rating"] - 1].value = a
    r[vh["result"] - 1].value = "PASS"

ws = wb["README"]
ws.cell(ws.max_row + 1, 1, "별점 분포")
ws.cell(ws.max_row, 2, "현실적인 J자형 분포로 조정 (5점 180 / 4점 55 / 3점 18 / 2점 10 / 1점 24). "
        "1~2점 리뷰는 review_question.review_answer 에 NEGATIVE 값 사용 (기존 POSITIVE/NEUTRAL 에 추가)")
wb.save(PATH)
from collections import Counter
print("완료:", dict(sorted(Counter(v for l in agg.values() for v in l).items())),
      "평균", round(sum(v for l in agg.values() for v in l) / sum(map(len, agg.values())), 2))
