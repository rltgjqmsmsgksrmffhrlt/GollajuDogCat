# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference
import pandas as pd

FONT_NAME = "Arial"
thin = Side(style="thin", color="B7B7B7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
body_font = Font(name=FONT_NAME, size=9)
title_font = Font(name=FONT_NAME, bold=True, size=13)
section_font = Font(name=FONT_NAME, bold=True, size=11, color="1F4E78")

wb = openpyxl.Workbook()

def style_header_row(ws, row, n_cols):
    for col in range(1, n_cols+1):
        c = ws.cell(row=row, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

def style_body(ws, start_row, end_row, n_cols):
    for r in range(start_row, end_row+1):
        for c in range(1, n_cols+1):
            cell = ws.cell(row=r, column=c)
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

# ===========================================================================
# 시트0: 개요
# ===========================================================================
ws0 = wb.active
ws0.title = "개요"
ws0["B2"] = "골라주개냥(우지빌) A/B 테스트 가상 운영 데이터 — 분석 요약"
ws0["B2"].font = title_font

overview_lines = [
    "",
    "생성 개요",
    "- 기간: 2026-09-08 ~ 2026-12-06 (90일)",
    "- 대상: A/B 각 2,500명 (총 5,000명), 90일간 순차 유입",
    "- 검증 가설 3개: 선택 피로도 해소 / 서비스 충성도 / 유저 데이터 확보",
    "- 이벤트 로그 약 27만 건, 주문 약 1.2만 건 생성",
    "",
    "가설 요약",
    "[가설1] 선택 피로도 해소: 추천 근거·점수 제공(A) 시 장바구니 이탈률↓, 구매전환율↑",
    "[가설2] 서비스 충성도: 고도화된 맞춤 추천·비교 경험이 서비스 충성도(NPS)를 높인다",
    "[가설3] 유저 데이터 확보: 반응 기록이 추천에 반영됨을 인지하면 반응 기록률이 높아진다",
    "",
    "시트 구성",
    "1. AB지표_검증결과 — 원본 문서 기준치와 실제 생성 데이터 지표 비교",
    "2. 퍼널_구매여정 — 가설1 세션 단위 퍼널 (진입→탐색→장바구니→구매)",
    "3. 퍼널_반응입력여정 — 가설3 주문 단위 퍼널 (알람→확인→시작→완료→재진입)",
    "4. 코호트_재구매율 — 가입 주차별 코호트의 주차별 구매 유저 비율(%)",
    "",
    "※ 참고: '장바구니 이탈률(A 목표 40%)'와 '추천상품 CVR(A 목표 3.3%)'는",
    "   원본 문서상 서로 다른 벤치마크 출처(전자상거래 일반 이탈률 vs 반려동물 카테고리 CVR)를",
    "   사용하고 있어 동일 세션 모집단에서 두 값을 동시에 정확히 재현하는 것이 수학적으로 불가능했습니다.",
    "   본 데이터는 추천상품 CVR·비교 후 전환율 등 세부 지표를 우선 재현하고,",
    "   장바구니 이탈률은 그 결과로 나온 근사치(A 56.6%)를 사용했습니다.",
]
for i, line in enumerate(overview_lines, start=3):
    c = ws0.cell(row=i, column=2, value=line)
    if line.startswith("[") or (line and not line.startswith("-") and not line.startswith("※") and not line.startswith(" ") and i > 3 and line not in ("",)):
        if line in ("생성 개요", "가설 요약", "시트 구성"):
            c.font = section_font
    c.font = Font(name=FONT_NAME, size=10, bold=(line in ("생성 개요","가설 요약","시트 구성")))
ws0.column_dimensions["B"].width = 100

# ===========================================================================
# 시트1: AB지표 검증결과
# ===========================================================================
ws1 = wb.create_sheet("AB지표_검증결과")
headers1 = ["가설", "지표", "그룹", "원본 문서 목표치", "생성 데이터 실측치", "비고"]
for col, h in enumerate(headers1, start=1):
    ws1.cell(row=1, column=col, value=h)
style_header_row(ws1, 1, len(headers1))

metric_rows = [
    ("가설1: 선택 피로도 해소", "장바구니 이탈률", "A", "40%", "56.6%", "추천CVR·비교전환율과 동일 모집단에서 완전 양립 불가, 근사치"),
    ("가설1: 선택 피로도 해소", "장바구니 이탈률", "B", "52.8~56.1%", "54.4%", "목표 범위 내"),
    ("가설1: 선택 피로도 해소", "추천 상품 CVR", "A", "3.3%", "3.73%", "근접"),
    ("가설1: 선택 피로도 해소", "추천 상품 CVR", "B", "3.21%", "3.09%", "근접"),
    ("가설1: 선택 피로도 해소", "상품상세 체류시간", "A", "60초", "59.9초", "정확히 일치"),
    ("가설1: 선택 피로도 해소", "상품상세 체류시간", "B", "43초", "43.1초", "정확히 일치"),
    ("가설1: 선택 피로도 해소", "구매결정 소요시간", "A", "2일", "1.76일", "근접"),
    ("가설1: 선택 피로도 해소", "구매결정 소요시간", "B", "2.4~3.3일", "2.47일", "목표 범위 내"),
    ("가설2: 서비스 충성도", "NPS", "A", "55점 이상", "53.8점", "근접"),
    ("가설2: 서비스 충성도", "NPS", "B", "20~30점", "25.1점", "목표 범위 내"),
    ("가설2: 서비스 충성도(순수달성)", "맞춤추천 CVR", "A", "초기3.3%/목표5%", "5.75%", "목표 상회"),
    ("가설2: 서비스 충성도(순수달성)", "맞춤추천 CVR", "B", "-", "4.51%", "참고치"),
    ("가설2: 서비스 충성도(순수달성)", "상품 반응 데이터 입력률", "A", "초기5%/목표10%", "28.2%", "목표 상회(리뷰+반응 합산 특성상 높게 산출)"),
    ("가설2: 서비스 충성도(순수달성)", "상품 반응 데이터 입력률", "B", "-", "16.3%", "참고치"),
    ("가설2: 서비스 충성도(순수달성)", "N일 재구매율(30일)", "A", "8~10%", "11.4%", "근접"),
    ("가설2: 서비스 충성도(순수달성)", "N일 재구매율(30일)", "B", "-", "9.4%", "참고치"),
    ("가설2: 서비스 충성도(순수달성)", "N일 재구매율(60일)", "A", "12~15%", "13.3%", "목표 범위 내(표본 적음, n=300)"),
    ("가설2: 서비스 충성도(순수달성)", "N일 재구매율(60일)", "B", "-", "16.6%", "참고치(표본 적음, n=326)"),
    ("가설2: 서비스 충성도(순수달성)", "PDP경유 비교기능 사용률", "A", "15~25%", "22.2%", "목표 범위 내"),
    ("가설2: 서비스 충성도(순수달성)", "PDP경유 비교기능 사용률", "B", "-", "10.1%", "참고치"),
    ("가설2: 서비스 충성도(순수달성)", "전체 비교기능 사용률", "A", "10~20%", "26.3%", "목표 상회"),
    ("가설2: 서비스 충성도(순수달성)", "전체 비교기능 사용률", "B", "-", "12.5%", "참고치"),
    ("가설2: 서비스 충성도(순수달성)", "비교 후 전환율", "A", "4~6%", "6.6%", "근접(약간 상회)"),
    ("가설2: 서비스 충성도(순수달성)", "비교 후 전환율", "B", "-", "3.6%", "참고치"),
    ("가설3: 유저 데이터 확보", "반응입력 시작률", "A", "20~30%", "28.4%", "목표 범위 내"),
    ("가설3: 유저 데이터 확보", "반응입력 시작률", "B", "20% 이하", "18.7%", "목표 범위 내"),
    ("가설3: 유저 데이터 확보", "반응입력 완료율", "A", "85% 이상", "86.7%", "목표 범위 내"),
    ("가설3: 유저 데이터 확보", "반응입력 완료율", "B", "80% 이하", "76.5%", "목표 범위 내"),
    ("가설3: 유저 데이터 확보", "퍼널 이탈률", "A", "15% 이하", "13.3%", "목표 범위 내"),
    ("가설3: 유저 데이터 확보", "퍼널 이탈률", "B", "약 20%", "23.6%", "근접"),
    ("가설3: 유저 데이터 확보", "맞춤추천 재진입률", "A", "5% 이상", "5.4%", "목표 범위 내"),
    ("가설3: 유저 데이터 확보", "맞춤추천 재진입률", "B", "A보다 낮음", "2.0%", "목표 조건(A>B) 충족"),
    ("가설3: 유저 데이터 확보(순수달성)", "유효 반응 데이터 확보율", "A", "초기10%/목표15%", "22.8%", "목표 상회"),
    ("가설3: 유저 데이터 확보(순수달성)", "유효 반응 데이터 확보율", "B", "-", "12.4%", "참고치"),
]
for r_idx, row in enumerate(metric_rows, start=2):
    for c_idx, val in enumerate(row, start=1):
        ws1.cell(row=r_idx, column=c_idx, value=val)
style_body(ws1, 2, 1+len(metric_rows), len(headers1))

# merge 가설 column
def merge_col(ws, col, first_row, last_row):
    r = first_row
    while r <= last_row:
        v = ws.cell(row=r, column=col).value
        r2 = r
        while r2+1 <= last_row and ws.cell(row=r2+1, column=col).value == v:
            r2 += 1
        if r2 > r:
            ws.merge_cells(start_row=r, start_column=col, end_row=r2, end_column=col)
        r = r2 + 1
merge_col(ws1, 1, 2, 1+len(metric_rows))
merge_col(ws1, 2, 2, 1+len(metric_rows))

widths1 = {1: 28, 2: 22, 3: 6, 4: 18, 5: 16, 6: 40}
for col, w in widths1.items():
    ws1.column_dimensions[get_column_letter(col)].width = w
ws1.freeze_panes = "A2"

# ===========================================================================
# 시트2: 퍼널_구매여정
# ===========================================================================
funnel1 = pd.read_csv("../data/generated/funnel1_purchase_journey.csv")
ws2 = wb.create_sheet("퍼널_구매여정")
for col, h in enumerate(funnel1.columns, start=1):
    ws2.cell(row=1, column=col, value=h)
style_header_row(ws2, 1, len(funnel1.columns))
for r_idx, row in enumerate(funnel1.itertuples(index=False), start=2):
    for c_idx, val in enumerate(row, start=1):
        ws2.cell(row=r_idx, column=c_idx, value=val)
style_body(ws2, 2, 1+len(funnel1), len(funnel1.columns))
for col in range(1, len(funnel1.columns)+1):
    ws2.column_dimensions[get_column_letter(col)].width = 18
ws2.column_dimensions["A"].width = 22

# 퍼널 차트
chart1 = BarChart()
chart1.title = "구매 여정 퍼널 (세션 수)"
chart1.y_axis.title = "세션 수"
data = Reference(ws2, min_col=2, max_col=3, min_row=1, max_row=1+len(funnel1))
cats = Reference(ws2, min_col=1, min_row=2, max_row=1+len(funnel1))
chart1.add_data(data, titles_from_data=True)
chart1.set_categories(cats)
ws2.add_chart(chart1, "H2")

# ===========================================================================
# 시트3: 퍼널_반응입력여정
# ===========================================================================
funnel2 = pd.read_csv("../data/generated/funnel2_reaction_journey.csv")
ws3 = wb.create_sheet("퍼널_반응입력여정")
for col, h in enumerate(funnel2.columns, start=1):
    ws3.cell(row=1, column=col, value=h)
style_header_row(ws3, 1, len(funnel2.columns))
for r_idx, row in enumerate(funnel2.itertuples(index=False), start=2):
    for c_idx, val in enumerate(row, start=1):
        ws3.cell(row=r_idx, column=c_idx, value=val)
style_body(ws3, 2, 1+len(funnel2), len(funnel2.columns))
for col in range(1, len(funnel2.columns)+1):
    ws3.column_dimensions[get_column_letter(col)].width = 18
ws3.column_dimensions["A"].width = 18

chart2 = BarChart()
chart2.title = "반응입력 여정 퍼널 (주문 수)"
chart2.y_axis.title = "주문 수"
data2 = Reference(ws3, min_col=2, max_col=3, min_row=1, max_row=1+len(funnel2))
cats2 = Reference(ws3, min_col=1, min_row=2, max_row=1+len(funnel2))
chart2.add_data(data2, titles_from_data=True)
chart2.set_categories(cats2)
ws3.add_chart(chart2, "H2")

# ===========================================================================
# 시트4: 코호트_재구매율
# ===========================================================================
cohort_A = pd.read_csv("../data/generated/cohort_A_repurchase_pct.csv")
cohort_B = pd.read_csv("../data/generated/cohort_B_repurchase_pct.csv")

ws4 = wb.create_sheet("코호트_재구매율")
ws4.cell(row=1, column=1, value="A그룹 — 가입 주차별 코호트 x 경과 주차별 구매 유저 비율(%)")
ws4["A1"].font = section_font
for col, h in enumerate(cohort_A.columns, start=1):
    ws4.cell(row=2, column=col, value=h if h != "signup_week" else "가입주차")
style_header_row(ws4, 2, len(cohort_A.columns))
for r_idx, row in enumerate(cohort_A.itertuples(index=False), start=3):
    for c_idx, val in enumerate(row, start=1):
        ws4.cell(row=r_idx, column=c_idx, value=val)
style_body(ws4, 3, 2+len(cohort_A), len(cohort_A.columns))

start_row_b = 2 + len(cohort_A) + 3
ws4.cell(row=start_row_b, column=1, value="B그룹 — 가입 주차별 코호트 x 경과 주차별 구매 유저 비율(%)")
ws4.cell(row=start_row_b, column=1).font = section_font
for col, h in enumerate(cohort_B.columns, start=1):
    ws4.cell(row=start_row_b+1, column=col, value=h if h != "signup_week" else "가입주차")
style_header_row(ws4, start_row_b+1, len(cohort_B.columns))
for r_idx, row in enumerate(cohort_B.itertuples(index=False), start=start_row_b+2):
    for c_idx, val in enumerate(row, start=1):
        ws4.cell(row=r_idx, column=c_idx, value=val)
style_body(ws4, start_row_b+2, start_row_b+1+len(cohort_B), len(cohort_B.columns))

for col in range(1, len(cohort_A.columns)+1):
    ws4.column_dimensions[get_column_letter(col)].width = 12
ws4.column_dimensions["A"].width = 14

wb.save("../data/generated/AB테스트_가상데이터_분석요약.xlsx")
print("saved")
