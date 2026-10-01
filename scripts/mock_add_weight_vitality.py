"""review_question 시트에 WEIGHT_VITALITY(체중·활력) 문항을 리뷰당 1행씩 추가한다.

- 기존 5개 문항 뒤에 id를 이어서 추가 (리뷰 287건 -> 287행)
- 답변: 4~5점은 POSITIVE 85% / NEUTRAL 15%, 3점 이하는 NEUTRAL (1~2점 중 일부는 NEGATIVE)
- updated_at 은 같은 리뷰의 기존 문항 값을 그대로 사용. 시드 고정, 이미 추가돼 있으면 종료.
"""
import os
import numpy as np
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), "..", "data", "mock", "상품 목데이터",
                    "golajulgaenyang_full_schema_mockdata_product_thumbnails_linked.xlsx")
RNG = np.random.default_rng(2027)
wb = openpyxl.load_workbook(PATH)
qs, rv = wb["review_question"], wb["review"]
qh = {c.value: c.column for c in qs[1]}
rh = {c.value: c.column for c in rv[1]}
rows = [r for r in qs.iter_rows(min_row=2) if r[0].value is not None]
if any(r[qh["review_question_type"] - 1].value == "WEIGHT_VITALITY" for r in rows):
    print("이미 추가돼 있습니다."); raise SystemExit

upd = {r[qh["review_id"] - 1].value: r[qh["updated_at"] - 1].value for r in rows}
nid = max(r[qh["id"] - 1].value for r in rows) + 1
for r in rv.iter_rows(min_row=2):
    rid, star = r[rh["id"] - 1].value, r[rh["star_rate"] - 1].value
    if rid is None:
        continue
    if star >= 4:
        ans = "POSITIVE" if RNG.random() < 0.85 else "NEUTRAL"
    elif star == 3:
        ans = "NEUTRAL"
    else:
        ans = "NEGATIVE" if RNG.random() < 0.3 else "NEUTRAL"
    row = [None] * len(qh)
    row[qh["id"] - 1], row[qh["review_id"] - 1] = nid, rid
    row[qh["review_question_type"] - 1], row[qh["review_answer"] - 1] = "WEIGHT_VITALITY", ans
    row[qh["updated_at"] - 1] = upd[rid]
    qs.append(row); nid += 1

ws = wb["README"]
ws.cell(ws.max_row + 1, 1, "리뷰 질문 유형")
ws.cell(ws.max_row, 2, "리뷰당 6개 구조화 질문 응답 (PALATABILITY/DIGESTION/FEEDING_CONVENIENCE/SKIN_COAT/WEIGHT_VITALITY/ALLERGY)")
wb.save(PATH)
print("추가 완료:", nid - 1436, "행")
